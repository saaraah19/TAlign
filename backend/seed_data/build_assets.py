"""
Renders the demo resumes and knowledge documents to real PDF / DOCX / TXT
files. The rendered files are committed to the repository, so you only need
to run this if you edit resume_content.py or knowledge_content.py:

    pip install reportlab python-docx        # python-docx is already a backend dependency
    python seed_data/build_assets.py

Output: seed_data/resumes/*.{pdf,docx,txt} and seed_data/knowledge/*.{pdf,docx,txt}.
"""

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from knowledge_content import DOCUMENTS
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate
from resume_content import RESUMES, SCANNED_RESUME

HERE = Path(__file__).parent
RESUME_DIR = HERE / "resumes"
KNOWLEDGE_DIR = HERE / "knowledge"

_styles = getSampleStyleSheet()
H1 = ParagraphStyle("H1", parent=_styles["Title"], fontSize=20, spaceAfter=2, alignment=0)
SUB = ParagraphStyle("Sub", parent=_styles["Normal"], fontSize=11, textColor="#444444", spaceAfter=2)
H2 = ParagraphStyle("H2", parent=_styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=3)
BODY = ParagraphStyle("Body", parent=_styles["Normal"], fontSize=10, leading=14, spaceAfter=3)
BULLET = ParagraphStyle("Bullet", parent=BODY, leftIndent=14, bulletIndent=4)


def _p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(text), style)


# --------------------------------------------------------------------- resumes
def resume_blocks(r: dict) -> list[tuple[str, str]]:
    """The resume as (kind, text) pairs, shared by every output format."""
    blocks = [("title", r["name"]), ("sub", r["headline"]), ("sub", r["contact"]),
              ("h2", "Summary"), ("body", r["summary"]), ("h2", "Experience")]
    for title, company, dates, bullets in r["experience"]:
        blocks.append(("body", f"{title} — {company} ({dates})"))
        blocks += [("bullet", b) for b in bullets]
    blocks += [("h2", "Skills"), ("body", r["skills"]), ("h2", "Education"), ("body", r["education"])]
    return blocks


def write_pdf(path: Path, blocks: list[tuple[str, str]]) -> None:
    style = {"title": H1, "sub": SUB, "h2": H2, "body": BODY, "bullet": BULLET}
    story = []
    for kind, text in blocks:
        if kind == "bullet":
            story.append(Paragraph(escape(text), BULLET, bulletText="•"))
        else:
            story.append(_p(text, style[kind]))
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=1.8 * cm, bottomMargin=1.8 * cm, title=blocks[0][1]).build(story)


def write_docx(path: Path, blocks: list[tuple[str, str]]) -> None:
    doc = Document()
    for kind, text in blocks:
        if kind == "title":
            doc.add_heading(text, level=0)
        elif kind == "h2":
            doc.add_heading(text, level=2)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    doc.save(str(path))


def write_txt(path: Path, blocks: list[tuple[str, str]]) -> None:
    lines: list[str] = []
    for kind, text in blocks:
        if kind == "title":
            lines += [text.upper(), ""]
        elif kind == "h2":
            lines += ["", text.upper(), "-" * len(text)]
        elif kind == "bullet":
            lines.append(f"  - {text}")
        else:
            lines.append(text)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


WRITERS = {"pdf": write_pdf, "docx": write_docx, "txt": write_txt}


def write_scanned_pdf(path: Path) -> None:
    """A PDF that looks like a page but holds only drawn shapes: no text to extract."""
    c = canvas.Canvas(str(path), pagesize=A4)
    _, height = A4
    c.setFillGray(0.85)
    c.rect(2 * cm, height - 4 * cm, 8 * cm, 1 * cm, fill=1, stroke=0)  # "name" bar
    for i in range(22):
        y = height - (6 + i * 0.9) * cm
        c.rect(2 * cm, y, (14 + (i % 4)) * cm, 0.35 * cm, fill=1, stroke=0)  # "text lines"
    c.save()


def main() -> None:
    RESUME_DIR.mkdir(exist_ok=True)
    KNOWLEDGE_DIR.mkdir(exist_ok=True)
    for key, r in RESUMES.items():
        WRITERS[r["format"]](RESUME_DIR / f"{key}.{r['format']}", resume_blocks(r))
    write_scanned_pdf(RESUME_DIR / f"{SCANNED_RESUME}.pdf")

    for key, d in DOCUMENTS.items():
        blocks = [("title", d["title"])]
        for heading, paragraphs in d["sections"]:
            blocks.append(("h2", heading))
            blocks += [("body", p) for p in paragraphs]
        WRITERS[d["format"]](KNOWLEDGE_DIR / f"{key}.{d['format']}", blocks)
    print(f"wrote {len(list(RESUME_DIR.iterdir()))} resumes and {len(list(KNOWLEDGE_DIR.iterdir()))} documents")


if __name__ == "__main__":
    main()
